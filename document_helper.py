from db import db, Negocio, Cartelera, Documento 
from flask import  jsonify, request,   session
import os
from datetime import datetime
from werkzeug.utils import secure_filename
from docx import Document
from pdf2image import convert_from_path
from PIL import Image, ImageDraw, ImageFont
import openpyxl
from helper import logger

user_folder = None
#filename = None

def manejar_post(id, app):
    logger.info("Se ha recibido una solicitud POST.")
    
    if 'ruta_imagen' not in request.files:
        return error_respuesta("No se ha seleccionado ningún archivo.")
    
    file = request.files['ruta_imagen']
    if file.filename == '':
        return error_respuesta("El archivo seleccionado está vacío.")
    
    tipo_anuncio = request.form.get('tipo_anuncio')
    if not tipo_anuncio:
        return error_respuesta("Tipo de anuncio no especificado.")
    
    usuario_id = session.get('user_id')
    usuario_nombre = session.get('usuario')
    user_folder = crear_directorio(usuario_nombre, usuario_id, id, app.config['UPLOAD_FOLDER'])
    
    file_path = guardar_archivo(file, user_folder)
    if not file_path:
        return error_respuesta("Error al guardar el archivo.")

    return procesar_documento(file_path, tipo_anuncio, id)

def error_respuesta(mensaje):
    logger.info(mensaje)
    return jsonify({'success': False, 'message': mensaje}), 400

def crear_directorio(usuario_nombre, usuario_id, id, folder):
    folder_name = f"{usuario_nombre}_{usuario_id}_{id}"
    user_folder = os.path.join(folder, folder_name)
    os.makedirs(user_folder, exist_ok=True)
    return user_folder

def guardar_archivo(file, user_folder):
    filename = secure_filename(file.filename)
    file_path = os.path.join(user_folder, filename)
    logger.info(f"Guardando archivo en: {file_path}")
    file.save(file_path)
    return file_path

def procesar_documento(file_path, tipo_anuncio, id):
    filename = os.path.basename(file_path)
    
    if filename.lower().endswith(('.jpg', '.jpeg', '.png', '.gif')):
        return guardar_imagen(file_path, tipo_anuncio, id)
    
    elif filename.lower().endswith('.pdf'):
        return procesar_pdf(file_path, tipo_anuncio, id, filename)
    
    elif filename.lower().endswith(('.doc', '.docx')):
        return procesar_doc(file_path, tipo_anuncio, id, filename)
    
    elif filename.lower().endswith('.xlsx'):
        return procesar_xlsx(file_path, tipo_anuncio, id, filename)
    
    return error_respuesta("Tipo de archivo no soportado. Solo se permiten imágenes, PDF, DOC o XLSX.")

def guardar_imagen(file_path, tipo_anuncio, id):
    logger.info("Guardando una imagen.")
    nuevo_documento = Documento(
        id_cartelera=id,
        tipo_documento=tipo_anuncio,
        archivo=file_path,
        fecha_subida=datetime.now()
    )
    db.session.add(nuevo_documento)
    db.session.commit()
    logger.info(f"Documento guardado: {file_path}")
    return jsonify({'success': True, 'message': 'Documentos convertidos y guardados con éxito!'}), 200

def procesar_pdf(file_path, tipo_anuncio, id, filename):
    logger.info(f"Procesando un archivo PDF. {file_path}")
    images = convert_from_path(file_path)
    return guardar_imagenes_de_pdf(images, tipo_anuncio, id, file_path, filename)

def guardar_imagenes_de_pdf(images, tipo_anuncio, id, file_path, filename):
    logger.info("Guardando imagenes de un PDF.")
    logger.info(f"File path: {file_path} - Tamanio: {len(images)} - Tipo: {type(images)} , Tipo de anuncio {tipo_anuncio}")
    user_folder = os.path.dirname(file_path)
    for i, image in enumerate(images):
        image_filename = f"{filename.split('.')[0]}_{i}.png"
        image_path = os.path.join(user_folder, image_filename)
        image.save(image_path, 'PNG')
        guardar_documento(image_path, tipo_anuncio, id)
    
    return jsonify({'success': True, 'message': 'Documentos convertidos y guardados con éxito!'}), 200

def procesar_doc(file_path, tipo_anuncio, id, filename):
    logger.info("Procesando un archivo DOC/DOCX.")
    doc = Document(file_path)
    img = convertir_doc_a_imagen(doc, filename)
    guardar_documento(img, tipo_anuncio, id)
    return jsonify({'success': True, 'message': 'Documentos convertidos y guardados con éxito!'}), 200

def convertir_doc_a_imagen(doc, filename):
    img_width = 800
    img_height = len(doc.paragraphs) * 20 + 20
    img = Image.new('RGB', (img_width, img_height), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    
    try:
        font = ImageFont.truetype("arial.ttf", 16)
    except IOError:
        font = ImageFont.load_default()

    y_text = 10
    for para in doc.paragraphs:
        draw.text((10, y_text), para.text, fill=(0, 0, 0), font=font)
        y_text += 20

    image_filename = f"{filename.split('.')[0]}.png"
    image_path = os.path.join(user_folder, image_filename)
    img.save(image_path)
    return image_path

def procesar_xlsx(file_path, tipo_anuncio, id, filename):
    logger.info("Procesando un archivo XLSX.")
    wb = openpyxl.load_workbook(file_path)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        img = convertir_xlsx_a_imagen(ws, filename, sheet)
        guardar_documento(img, tipo_anuncio, id)
    
    return jsonify({'success': True, 'message': 'Documentos convertidos y guardados con éxito!'}), 200

def convertir_xlsx_a_imagen(ws, filename, sheet):
    img = Image.new('RGB', (800, 400), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    try:
        font = ImageFont.truetype("arial.ttf", 16)
    except IOError:
        font = ImageFont.load_default()

    y_text = 10
    for row in ws.iter_rows(values_only=True):
        line = " | ".join(str(cell) for cell in row if cell is not None)
        draw.text((10, y_text), line, fill=(0, 0, 0), font=font)
        y_text += 20

    image_filename = f"{filename.split('.')[0]}_{sheet}.png"
    image_path = os.path.join(user_folder, image_filename)
    img.save(image_path)
    return image_path

def guardar_documento(image_path, tipo_anuncio, id):
    nuevo_documento = Documento(
        id_cartelera=id,
        tipo_documento=tipo_anuncio,
        archivo=image_path,
        fecha_subida=datetime.now()
    )
    db.session.add(nuevo_documento)
    db.session.commit()
    logger.info(f"Documento guardado: {image_path}")